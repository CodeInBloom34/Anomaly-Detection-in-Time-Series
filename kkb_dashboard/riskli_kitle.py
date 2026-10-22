"""KKB trendlerinden riskli kitle tespiti.

Soru: Hangi musteri kendi normalinden olumsuz yonde sapiyor ve bu sapma gecmiste
gercekten gecikme / takip ile sonuclanan desenlere benziyor?

Adimlar:
1. Her sorguda, musterinin o ana kadarki gecmisinden bir "normal" (medyan) cikarilir.
2. Son deger bu normalle karsilastirilir; yalnizca bankacilik acisindan OLUMSUZ yondeki
   sapmalar sayilir (doluluk artisi, not dususu, banka sayisi artisi, limit kisilmasi...).
3. Gecmis sorgular uzerinde "sonraki 12 ayda gecikme veya takip" hedefiyle aciklanabilir
   bir lojistik regresyon egitilir. Model, hangi sapma desenlerinin gercekten kotu
   sonuclandigini ogrenir ve her musteriye bir olasilik verir.
4. Model, zamana gore ayrilmis bir test doneminde dogrulanir (gecmisle egit, sonrayla test).
5. Su an gecikmede olmayan musteriler son sorgularina gore puanlanir; olasiligi portfoy
   ortalamasinin belirgin uzerinde olanlar riskli kitle listesine girer.

Kullanim (kkb_dashboard klasorunden):
    python riskli_kitle.py --data data --out docs/Riskli_Kitle.csv
"""

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import roc_auc_score
from sklearn.preprocessing import StandardScaler

from anomaly_score import olcek

# metrik: (gorunen ad, olumsuz yon (+1 artis kotu, -1 azalis kotu), bicim)
METRIKLER = {
    "Doluluk": ("Limit doluluğu", +1, "yuzde"),
    "BankaSayisi": ("Çalışılan banka sayısı", +1, "adet"),
    "KKBNot": ("KKB notu", -1, "adet"),
    "RotatifPay": ("Rotatif kredi payı", +1, "yuzde"),
    "LogLimit": ("Bankaların toplam limiti", -1, "limit"),
    "LogRisk": ("Toplam sektör borcu", +1, "risk"),
    "CekKarsiliksizAdet": ("Karşılıksız çek", +1, "adet"),
}
SEVIYE = ["Doluluk", "KKBNot"]  # sapmanin yaninda mevcut seviye de bilgi tasir
UFUK_GUN = 365  # "sonraki" sorgu bu sure icindeyse etiket gecerli
TEST_BASLANGIC = pd.Timestamp("2026-04-01")
TAVAN = 6.0  # olumsuz sapma ust siniri (z birimi)
# Kademe: olasiligin portfoy ortalamasina orani (kat). Pilot sonuclarina gore ayarlanmali.
KADEME_KAT = {"Yüksek": 10, "Orta": 5, "İzleme": 3}
AKSIYON = {
    "Yüksek": "1 hafta içinde ilişki yöneticisi ziyareti; limit artışı dondurulur; teminat gözden geçirilir",
    "Orta": "2 hafta içinde analist incelemesi; yeni limit kararı kredi komitesine",
    "İzleme": "Aylık izleme listesine alınır; bir sonraki KKB sorgusu öne çekilir",
}


def hazirla(sorgu, detay):
    s = sorgu.sort_values(["VKN", "SorguTarihi"]).reset_index(drop=True).copy()
    rot = (detay[detay.KrediTuru == "Rotatif"].groupby("SorguID").Risk.sum()
           / detay.groupby("SorguID").Risk.sum()).rename("RotatifPay")
    s = s.join(rot, on="SorguID")
    s["RotatifPay"] = s.RotatifPay.fillna(0)
    s["Doluluk"] = s.SektorToplamRisk / s.SektorToplamLimit
    s["LogLimit"] = np.log(s.SektorToplamLimit.clip(lower=1))
    s["LogRisk"] = np.log(s.SektorToplamRisk.clip(lower=1))
    g = s.groupby("VKN")
    s["SonrakiTarih"] = g.SorguTarihi.shift(-1)
    s["SonrakiGecikme"] = g.GecikmeVar.shift(-1)
    s["SonrakiTakip"] = g.TakipVar.shift(-1)
    return s


def sapmalar(s):
    """Her sorgu icin: musterinin onceki sorgularinin medyani (normal) ve olumsuz yondeki sapma."""
    # Olcek tabani: portfoydeki ardisik degisimlerin tipik buyuklugu.
    taban = {m: olcek(s.groupby("VKN")[m].diff().dropna().to_numpy(dtype=float)) for m in METRIKLER}
    normal = pd.DataFrame(np.nan, index=s.index, columns=list(METRIKLER))
    for _, grp in s.groupby("VKN"):
        vals = grp[list(METRIKLER)].to_numpy(dtype=float)
        for i in range(1, len(grp)):
            normal.iloc[normal.index.get_loc(grp.index[i])] = np.median(vals[:i], axis=0)
    out = pd.DataFrame(index=s.index)
    for m, (_, yon, _) in METRIKLER.items():
        z = (s[m] - normal[m]) / taban[m]
        out[f"normal_{m}"] = normal[m]
        out[f"sapma_{m}"] = (yon * z).clip(lower=0, upper=TAVAN)
    return out


def ozellik_listesi():
    return [f"sapma_{m}" for m in METRIKLER] + [f"seviye_{m}" for m in SEVIYE]


def veri_seti(sorgu, detay):
    s = hazirla(sorgu, detay)
    s = s.join(sapmalar(s))
    for m in SEVIYE:
        s[f"seviye_{m}"] = s[m]
    temiz = (s.GecikmeVar == 0) & (s.TakipVar == 0)
    gecmisli = s.OncekiSorguID.notna()
    etiketli = s.SonrakiTarih.notna() & ((s.SonrakiTarih - s.SorguTarihi).dt.days <= UFUK_GUN)
    s["Hedef"] = ((s.SonrakiGecikme == 1) | (s.SonrakiTakip == 1)).astype(int)
    s["Egitime_Uygun"] = temiz & gecmisli & etiketli
    s["Puanlanabilir"] = temiz & gecmisli
    return s


class Model:
    """Lojistik regresyon; olumsuz sapma sinyalleri yalnizca riski ARTIRABILIR.

    Iliskili sinyaller yuzunden bir sapmanin agirligi negatif cikarsa ("limit kisildi, risk
    azaldi" gibi) bu sinyal cikarilip model yeniden kurulur. Boylece her neden, analistin
    okuyabilecegi yonde calisir.
    """

    def __init__(self, X, y):
        self.ozellikler = list(X.columns)
        while True:
            self.olc = StandardScaler().fit(X[self.ozellikler])
            self.lr = LogisticRegression(C=0.5, max_iter=2000).fit(self.olc.transform(X[self.ozellikler]), y)
            ters = [o for o, k in zip(self.ozellikler, self.lr.coef_[0]) if o.startswith("sapma_") and k < 0]
            if not ters:
                break
            self.ozellikler.remove(min(ters, key=lambda o: self.lr.coef_[0][self.ozellikler.index(o)]))

    def olasilik(self, X):
        return self.lr.predict_proba(self.olc.transform(X[self.ozellikler]))[:, 1]

    def katkilar(self, X):
        return pd.DataFrame(self.olc.transform(X[self.ozellikler]) * self.lr.coef_[0], columns=self.ozellikler,
                            index=X.index)

    @property
    def agirlik(self):
        return pd.Series(self.lr.coef_[0], index=self.ozellikler)


def dogrula(s):
    """Zamana gore ayrilmis test: gecmisle egit, sonraki donemde dene."""
    oz = ozellik_listesi()
    e = s[s.Egitime_Uygun]
    tr, te = e[e.SorguTarihi < TEST_BASLANGIC], e[e.SorguTarihi >= TEST_BASLANGIC]
    p = Model(tr[oz], tr.Hedef).olasilik(te[oz])
    taban_oran = te.Hedef.mean()
    ust = te.assign(p=p).nlargest(max(1, int(len(te) * 0.1)), "p")
    return {
        "egitim_satir": len(tr), "test_satir": len(te), "test_kotu": int(te.Hedef.sum()),
        "auc": roc_auc_score(te.Hedef, p) if te.Hedef.nunique() > 1 else np.nan,
        "taban_oran": taban_oran, "ust10_oran": ust.Hedef.mean(),
        "kaldirac": ust.Hedef.mean() / taban_oran if taban_oran else np.nan,
        "ust10_yakalama": ust.Hedef.sum() / max(te.Hedef.sum(), 1),
    }


def bicimle(deger, tur):
    if pd.isna(deger):
        return "-"
    if tur == "yuzde":
        return "%" + f"{deger * 100:.0f}"
    if tur in ("limit", "risk"):
        tl = np.exp(deger)
        for esik, ek in ((1e9, " milyar TL"), (1e6, " milyon TL")):
            if tl >= esik:
                return f"{tl / esik:.1f}".replace(".", ",") + ek
        return f"{tl:,.0f} TL".replace(",", ".")
    return f"{deger:.0f}"


def nedenler(satir, katkilar, n=2):
    sec = [(k, v) for k, v in katkilar.items() if k.startswith("sapma_") and v > 0]
    sec.sort(key=lambda kv: -kv[1])
    out = []
    for k, _ in sec[:n]:
        m = k.removeprefix("sapma_")
        ad, _, tur = METRIKLER[m]
        if satir[k] <= 0:
            continue
        out.append(f"{ad}: normali {bicimle(satir[f'normal_{m}'], tur)}, şimdi {bicimle(satir[m], tur)}")
    return out + [""] * (n - len(out))


def riskli_kitle(sorgu, detay, musteri):
    s = veri_seti(sorgu, detay)
    test = dogrula(s)
    oz = ozellik_listesi()
    e = s[s.Egitime_Uygun]
    m = Model(e[oz], e.Hedef)
    taban_oran = e.Hedef.mean()

    son = s[s.Puanlanabilir & (s.SonSorguMu == 1)].copy()
    son["Olasilik"] = m.olasilik(son)
    son["Kat"] = son.Olasilik / taban_oran
    son["Kademe"] = np.select([son.Kat >= KADEME_KAT["Yüksek"], son.Kat >= KADEME_KAT["Orta"],
                               son.Kat >= KADEME_KAT["İzleme"]], list(KADEME_KAT), "")
    katki = m.katkilar(son)
    ned = [nedenler(son.loc[i], katki.loc[i]) for i in son.index]
    son["Neden1"], son["Neden2"] = [x[0] for x in ned], [x[1] for x in ned]
    liste = son[son.Kademe != ""].merge(
        musteri[["VKN", "Unvan", "Segment", "NACEAciklama", "Bolge", "PortfoyYoneticisi"]], on="VKN", how="left")
    sira = {"Yüksek": 0, "Orta": 1, "İzleme": 2}
    liste = liste.sort_values(["Kademe", "Olasilik"], key=lambda c: c.map(sira) if c.name == "Kademe" else -c)

    bozuk = s[(s.SonSorguMu == 1) & ((s.GecikmeVar == 1) | (s.TakipVar == 1))].merge(
        musteri[["VKN", "Unvan", "Segment", "NACEAciklama", "Bolge", "PortfoyYoneticisi"]], on="VKN", how="left")
    return liste, bozuk, test, m.agirlik, taban_oran


SINYAL_ACIKLAMA = {
    "sapma_Doluluk": "Limit doluluğu müşterinin normaline göre arttı",
    "sapma_BankaSayisi": "Çalışılan banka sayısı müşterinin normaline göre arttı",
    "sapma_KKBNot": "KKB notu müşterinin normaline göre düştü",
    "sapma_RotatifPay": "Rotatif (işletme) kredisi payı arttı",
    "sapma_LogLimit": "Bankaların toplam limiti kısıldı",
    "sapma_LogRisk": "Toplam sektör borcu normalinden hızlı büyüdü",
    "sapma_CekKarsiliksizAdet": "Karşılıksız çek ortaya çıktı",
    "seviye_Doluluk": "Limit doluluğunun mevcut seviyesi (yüksekse risk artar)",
    "seviye_KKBNot": "KKB notunun mevcut seviyesi (yüksekse risk azalır)",
}


def excel_yaz(yol, liste, bozuk, test, agirlik, taban_oran):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    baslik_f = Font(name=F, bold=True, size=14)
    kalin = Font(name=F, bold=True, size=10)
    normal = Font(name=F, size=10)
    gri = Font(name=F, size=9, color="52514E")
    bas_dolgu = PatternFill("solid", fgColor="2A78D6")
    bas_font = Font(name=F, bold=True, size=10, color="FFFFFF")
    ince = Border(bottom=Side(style="thin", color="E1E0D9"))
    sar = Alignment(wrap_text=True, vertical="top")

    def baslik_satiri(ws, satir, basliklar, genislik):
        for j, (b, w) in enumerate(zip(basliklar, genislik), 1):
            c = ws.cell(satir, j, b)
            c.font, c.fill, c.alignment = bas_font, bas_dolgu, Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[satir].height = 30

    wb = Workbook()

    # --- Riskli Kitle -------------------------------------------------------------------
    rk = wb.active
    rk.title = "Riskli Kitle"
    bas = ["Sıra", "Müşteri", "VKN", "Segment", "Sektör", "Bölge", "Portföy Yön.", "Kademe",
           "Olumsuz sonuç olasılığı (12 ay)", "Portföy ortalamasının kaç katı", "Bizim risk (TL)",
           "Olasılıkla ağırlıklı risk (TL)", "Neden 1", "Neden 2", "Önerilen aksiyon", "Son KKB sorgusu"]
    baslik_satiri(rk, 1, bas, [6, 20, 13, 11, 22, 13, 12, 10, 14, 13, 16, 16, 46, 46, 50, 13])
    for i, r in enumerate(liste.itertuples(), 2):
        degerler = [i - 1, r.Unvan, r.VKN, r.Segment, r.NACEAciklama, r.Bolge, r.PortfoyYoneticisi, r.Kademe,
                    float(r.Olasilik), None, float(r.BizimRisk), None,
                    r.Neden1 or "", r.Neden2 or "", AKSIYON[r.Kademe], r.SorguTarihi.to_pydatetime()]
        for j, v in enumerate(degerler, 1):
            c = rk.cell(i, j, v)
            c.font, c.alignment, c.border = normal, sar, ince
        rk.cell(i, 10, f"=I{i}/'Yöntem'!$C$6").number_format = '0.0"x"'
        rk.cell(i, 12, f"=I{i}*K{i}").number_format = "#,##0"
        rk.cell(i, 9).number_format = "0.0%"
        rk.cell(i, 11).number_format = "#,##0"
        rk.cell(i, 16).number_format = "dd.mm.yyyy"
        for j in (10, 12):
            rk.cell(i, j).font, rk.cell(i, j).border = normal, ince
    son = len(liste) + 1
    for kademe, renk in (("Yüksek", "F8D7D7"), ("Orta", "FDE9CF"), ("İzleme", "FFF6D6")):
        rk.conditional_formatting.add(f"A2:P{son}", FormulaRule(formula=[f'$H2="{kademe}"'],
                                                               fill=PatternFill("solid", fgColor=renk)))
    rk.freeze_panes = "C2"
    rk.auto_filter.ref = f"A1:P{son}"

    # --- Ozet ---------------------------------------------------------------------------
    oz = wb.create_sheet("Özet", 0)
    oz.column_dimensions["A"].width = 26
    for col in "BCDE":
        oz.column_dimensions[col].width = 20
    oz["A1"] = "KKB trendlerine göre riskli kitle"
    oz["A1"].font = baslik_f
    oz["A2"] = (f"Puanlama tarihi: son KKB sorgusu ({liste.SorguTarihi.max():%d.%m.%Y}). Her müşteri kendi "
                "normaline göre değerlendirildi; yalnızca bankacılık açısından olumsuz yöndeki sapmalar sayıldı.")
    oz["A2"].font, oz["A2"].alignment = gri, Alignment(wrap_text=True)
    oz.merge_cells("A2:E2")
    oz.row_dimensions[2].height = 30
    baslik_satiri(oz, 4, ["Kademe", "Müşteri sayısı", "Bizim risk (TL)", "Olasılıkla ağırlıklı risk (TL)",
                          "Önerilen aksiyon"], [26, 16, 20, 24, 70])
    for i, k in enumerate(["Yüksek", "Orta", "İzleme"], 5):
        oz.cell(i, 1, k).font = kalin
        oz.cell(i, 2, f"=COUNTIF('Riskli Kitle'!$H:$H,A{i})")
        oz.cell(i, 3, f"=SUMIF('Riskli Kitle'!$H:$H,A{i},'Riskli Kitle'!$K:$K)").number_format = "#,##0"
        oz.cell(i, 4, f"=SUMIF('Riskli Kitle'!$H:$H,A{i},'Riskli Kitle'!$L:$L)").number_format = "#,##0"
        oz.cell(i, 5, AKSIYON[k])
        for j in range(1, 6):
            oz.cell(i, j).alignment = sar
            if j > 1:
                oz.cell(i, j).font = normal
    oz.cell(8, 1, "Toplam").font = kalin
    for j, col in ((2, "B"), (3, "C"), (4, "D")):
        c = oz.cell(8, j, f"=SUM({col}5:{col}7)")
        c.font = kalin
        if j > 2:
            c.number_format = "#,##0"
    oz.cell(10, 1, "Zaten gecikmede / takipte").font = kalin
    oz.cell(10, 2, "=COUNTA('Zaten Gecikmede'!B:B)-1")
    oz.cell(10, 3, "Bu müşteriler listede değil; ayrı sekmede. Riskli kitle, henüz bozulmamış müşterileri gösterir.")
    oz.cell(10, 3).font = gri
    oz.cell(12, 1, "Nasıl okunur").font = kalin
    for i, metin in enumerate([
        "Olumsuz sonuç olasılığı: müşterinin önümüzdeki 12 ayda başka bir bankada gecikmeye veya takibe düşme "
        "olasılığı tahmini.",
        "Kaç katı: bu olasılığın portföy ortalamasına oranı. Kademeler bu orana göre verilir "
        f"(Yüksek ≥ {KADEME_KAT['Yüksek']}x, Orta ≥ {KADEME_KAT['Orta']}x, İzleme ≥ {KADEME_KAT['İzleme']}x).",
        "Neden 1 ve 2: olasılığı en çok yükselten iki sapma; müşterinin normali ve bugünkü değeriyle.",
        "Olasılıkla ağırlıklı risk: olasılık × bizim risk. Aynı kademede önceliği bu sütuna göre verin.",
    ], 13):
        c = oz.cell(i, 1, "• " + metin)
        c.font, c.alignment = normal, Alignment(wrap_text=True, vertical="top")
        oz.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        oz.row_dimensions[i].height = 28

    # --- Zaten gecikmede ------------------------------------------------------------------
    zg = wb.create_sheet("Zaten Gecikmede")
    baslik_satiri(zg, 1, ["Sıra", "Müşteri", "VKN", "Segment", "Sektör", "Bölge", "Portföy Yön.", "Durum",
                          "Bizim risk (TL)", "Son KKB sorgusu"], [6, 20, 13, 11, 22, 13, 12, 12, 16, 13])
    bozuk = bozuk.sort_values("BizimRisk", ascending=False)
    for i, r in enumerate(bozuk.itertuples(), 2):
        durum = "Takipte" if r.TakipVar == 1 else "Gecikmede"
        for j, v in enumerate([i - 1, r.Unvan, r.VKN, r.Segment, r.NACEAciklama, r.Bolge, r.PortfoyYoneticisi,
                               durum, float(r.BizimRisk), r.SorguTarihi.to_pydatetime()], 1):
            c = zg.cell(i, j, v)
            c.font, c.border = normal, ince
        zg.cell(i, 9).number_format = "#,##0"
        zg.cell(i, 10).number_format = "dd.mm.yyyy"
    zg.freeze_panes = "C2"

    # --- Yontem ---------------------------------------------------------------------------
    yt = wb.create_sheet("Yöntem")
    yt.column_dimensions["A"].width = 4
    yt.column_dimensions["B"].width = 52
    yt.column_dimensions["C"].width = 18
    yt.column_dimensions["D"].width = 60
    yt["B1"] = "Yöntem ve doğrulama"
    yt["B1"].font = baslik_f
    yt["B3"] = "Model sonuçları (zamana göre ayrılmış test)"
    yt["B3"].font = kalin
    satirlar = [
        ("Eğitim dönemi gözlem sayısı", test["egitim_satir"], "0", f"{TEST_BASLANGIC:%d.%m.%Y} öncesi sorgular"),
        ("Test dönemi gözlem sayısı", test["test_satir"], "0", f"{TEST_BASLANGIC:%d.%m.%Y} ve sonrası"),
        ("Portföy ortalaması olumsuz sonuç oranı", float(taban_oran), "0.0%",
         "Tüm eğitim verisinde 12 ay içinde gecikme/takibe düşenlerin oranı; 'kaç katı' sütununun paydası"),
        ("Testte en riskli %10'un olumsuz sonuç oranı", float(test["ust10_oran"]), "0.0%", ""),
        ("Kaldıraç (en riskli %10 / test ortalaması)", float(test["kaldirac"]), '0.0"x"', ""),
        ("Testteki olumsuz sonuçların en riskli %10'da yakalanan payı", float(test["ust10_yakalama"]), "0%", ""),
        ("Ayırma gücü (AUC, 0,5 = rastgele, 1 = mükemmel)", float(test["auc"]), "0.00", ""),
    ]
    for i, (ad, deger, fmt, not_) in enumerate(satirlar, 4):
        yt.cell(i, 2, ad).font = normal
        c = yt.cell(i, 3, deger)
        c.number_format, c.font = fmt, normal
        yt.cell(i, 4, not_).font = gri
    yt["C6"].comment = Comment("Model çıktısı (riskli_kitle.py). 'Kaç katı' formülleri bu hücreye bağlıdır.", "Model")
    yt["B12"] = "Sinyallerin etkisi (standartlaştırılmış ağırlık; büyük = daha etkili)"
    yt["B12"].font = kalin
    for i, (oz_ad, w) in enumerate(agirlik.sort_values(key=abs, ascending=False).items(), 13):
        yt.cell(i, 2, SINYAL_ACIKLAMA.get(oz_ad, oz_ad)).font = normal
        c = yt.cell(i, 3, float(w))
        c.number_format, c.font = "0.00", normal
    r0 = 13 + len(agirlik) + 1
    yt.cell(r0, 2, "Adımlar").font = kalin
    adimlar = [
        "Her sorguda müşterinin o güne kadarki sorgularının medyanı 'normal' kabul edildi.",
        "Bugünkü değer normalle karşılaştırıldı; yalnızca olumsuz yöndeki sapmalar sayıldı "
        "(doluluk artışı, not düşüşü, banka sayısı artışı, limit kısılması, karşılıksız çek...).",
        "Geçmiş sorgularda hangi sapma desenlerinin sonraki 12 ayda gecikme veya takiple sonuçlandığı "
        "lojistik regresyonla öğrenildi. Olumsuz sapmalar yalnızca riski artırabilecek şekilde kısıtlandı.",
        "Model, geçmiş dönemle eğitilip sonraki dönemde test edildi (yukarıdaki tablo).",
        "Şu an gecikmede olmayan müşteriler son sorgularına göre puanlandı; ortalamanın belirgin üzerinde "
        "olanlar listeye alındı.",
    ]
    for i, a in enumerate(adimlar, r0 + 1):
        c = yt.cell(i, 2, f"{i - r0}. {a}")
        c.font, c.alignment = normal, Alignment(wrap_text=True, vertical="top")
        yt.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        yt.row_dimensions[i].height = 28
    r1 = r0 + len(adimlar) + 2
    yt.cell(r1, 2, "Dikkat").font = kalin
    for i, a in enumerate([
        "Bu dosyadaki veriler sentetiktir. Sentetik veride bozulma deseni net olduğu için test sonuçları "
        "gerçekte beklenenden yüksektir; gerçek veride performans ayrıca ölçülmelidir.",
        "Liste bir karar değil, bakış önceliğidir. Kredi kararı analistte kalır.",
        "Gerçek müşteri listesi paylaşılırken KVKK ve KKB kullanım koşullarına uyulmalı, yalnızca yetkili "
        "kişilerle paylaşılmalıdır.",
    ], r1 + 1):
        c = yt.cell(i, 2, "• " + a)
        c.font, c.alignment = normal, Alignment(wrap_text=True, vertical="top")
        yt.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        yt.row_dimensions[i].height = 28

    # Formuller Excel'de acilista hesaplansin (openpyxl onbellek deger yazamaz).
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(yol)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", default="data")
    ap.add_argument("--out", default="docs/Riskli_Kitle.csv")
    ap.add_argument("--xlsx", default="docs/Riskli_Kitle.xlsx")
    args = ap.parse_args()
    d = Path(args.data)
    sorgu = pd.read_csv(d / "Fact_Sorgu.csv", dtype={"VKN": str}, parse_dates=["SorguTarihi"])
    detay = pd.read_csv(d / "Fact_RiskDetay.csv")
    musteri = pd.read_csv(d / "Dim_Musteri.csv", dtype={"VKN": str})
    liste, bozuk, test, agirlik, taban = riskli_kitle(sorgu, detay, musteri)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    liste.to_csv(args.out, index=False, encoding="utf-8")
    if args.xlsx:
        excel_yaz(args.xlsx, liste, bozuk, test, agirlik, taban)
    print(f"Riskli kitle: {len(liste)} musteri ({liste.Kademe.value_counts().to_dict()}), "
          f"zaten gecikmede: {len(bozuk)}")
    print("Zaman bazli test:", {k: round(v, 3) if isinstance(v, float) else v for k, v in test.items()})
    print("Model agirliklari:", agirlik.round(2).to_dict())


if __name__ == "__main__":
    main()
